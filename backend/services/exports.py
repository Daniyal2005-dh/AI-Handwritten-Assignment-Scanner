"""Step 5 — the three output modes.

1. **Original-scan PDF**   page images only: looks exactly like the sheet.
2. **Searchable PDF**      same page images, plus an invisible text layer
                           aligned to each recognised line so search/copy work.
3. **Editable document**   DOCX (and plain TXT) reconstructed from the text.

User corrections (``Page.corrected_text``) always win over raw recognition in
the editable outputs. The searchable layer uses the per-line recognition
results, because line-level alignment is key to searchability.
"""

from __future__ import annotations

import io
from typing import Any

import numpy as np
from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

PAGE_W, PAGE_H = A4  # 595.27 x 841.89 pt


def _slug(text: str) -> str:
    keep = [c if c.isalnum() else "-" for c in (text or "assignment").strip()]
    slug = "".join(keep).strip("-").lower()
    return "-".join(s for s in slug.split("-") if s)[:60] or "assignment"


def _page_image(source: np.ndarray | None) -> np.ndarray:
    if source is None:
        return np.full((1000, 707), 255, dtype=np.uint8)
    return source


def _pdf_bytes(assignment, pages, draw_text_layer: bool) -> bytes:
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    margin = 10 * mm
    from reportlab.lib.utils import ImageReader

    for page in pages:
        img = _page_image(page.image_array)
        ih, iw = img.shape[:2]
        inner_w, inner_h = PAGE_W - 2 * margin, PAGE_H - 2 * margin
        scale = min(inner_w / iw, inner_h / ih)
        w, h = iw * scale, ih * scale
        x, y = (PAGE_W - w) / 2, (PAGE_H - h) / 2

        png = io.BytesIO()
        Image.fromarray(img).save(png, format="PNG")
        png.seek(0)
        c.drawImage(ImageReader(png), x, y, w, h)

        if draw_text_layer and page.lines():
            c.setFillColorRGB(0, 0, 0)
            for line in page.lines():
                text = (line.get("text") or "").strip()
                if not text:
                    continue
                lx, ly, lw, lh = line["bbox"]
                font_size = max(4.0, lh * scale * 0.8)
                c.setFont("Helvetica", font_size)
                # Invisible (renderMode 3) text placed at the line's position.
                c.drawString(
                    x + lx * scale,
                    y + h - (ly + lh * 0.8) * scale,
                    text,
                    mode=3,
                )
        c.showPage()

    c.save()
    return buffer.getvalue()


def export_original_pdf(assignment, pages) -> bytes:
    return _pdf_bytes(assignment, pages, draw_text_layer=False)


def export_searchable_pdf(assignment, pages) -> bytes:
    return _pdf_bytes(assignment, pages, draw_text_layer=True)


def export_docx(assignment, pages) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(assignment.name, level=1)
    if assignment.subject:
        doc.add_paragraph(assignment.subject)

    for page in pages:
        doc.add_heading(f"Page {page.index}", level=2)
        text = page.text
        blocks = [b for b in text.split("\n\n") if b.strip()] if text else [""]
        paragraphs_meta = {tuple(p["lines"]): p.get("heading") for p in page.paragraphs()}
        line_ids = [l["i"] for l in page.lines()]
        using_recognized = not page.corrected_text and bool(blocks)

        offset = 0
        for block in blocks:
            if not block.strip():
                continue
            is_heading = False
            if using_recognized and paragraphs_meta:
                consumed = block.count("\n") + 1
                meta_key = tuple(line_ids[offset : offset + consumed])
                is_heading = paragraphs_meta.get(meta_key, False)
                offset += consumed
            if is_heading:
                doc.add_heading(block, level=3)
            else:
                doc.add_paragraph(block)
        if page is not pages[-1]:
            doc.add_page_break()

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out.read()


def export_txt(assignment, pages) -> bytes:
    parts = []
    for page in pages:
        header = f"Assignment: {assignment.name}"
        if assignment.subject:
            header += f"\nSubject: {assignment.subject}"
        parts.append(f"{header}\n\n--- Page {page.index} ---\n\n{page.text}".strip())
    return ("\n\n".join(parts)).encode("utf-8")


EXPORT_KINDS: dict[str, dict[str, Any]] = {
    "original_pdf": {
        "builder": export_original_pdf,
        "media_type": "application/pdf",
        "ext": "pdf",
    },
    "searchable_pdf": {
        "builder": export_searchable_pdf,
        "media_type": "application/pdf",
        "ext": "pdf",
    },
    "docx": {
        "builder": export_docx,
        "media_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "ext": "docx",
    },
    "txt": {"builder": export_txt, "media_type": "text/plain; charset=utf-8", "ext": "txt"},
}
